"""Blank ASMIS tracker formatted as an Excel Table named MishapTracker (for the Power Automate kit).
Usage: python3 data-generation/build_table_template.py   (needs openpyxl). No data, headers only."""
import os, openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = openpyxl.load_workbook(os.path.join(ROOT, "dashboards", "safety", "asmis_template.xlsx")).active
headers = [c.value for c in src[1] if c.value]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tracker"
ws.append(headers); ws.append([None] * len(headers))
t = Table(displayName="MishapTracker", ref=f"A1:{get_column_letter(len(headers))}2")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(t)
for i, h in enumerate(headers, 1): ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 2)
ws.freeze_panes = "A2"
wb.properties.creator = "Safety Office"; wb.properties.lastModifiedBy = "Safety Office"
out = os.path.join(ROOT, "dashboards", "safety", "asmis_tracker_table.xlsx"); wb.save(out)
print(out, len(headers), "columns, table MishapTracker")
