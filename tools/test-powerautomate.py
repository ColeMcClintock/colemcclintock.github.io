#!/usr/bin/env python3
"""Simulates the Power Automate auto-refresh flow end to end with FICTIONAL data and checks the result in headless Chrome.

  1. Generates fake FY trackers + a roll-up (data-generation/generate_asmis_fy_trackers.py) and formats each
     sheet as an Excel Table named MishapTracker.
  2. Converts each table to the JSON Excel Online "List rows present in a table" returns: every value a string,
     '@odata.etag' + 'ItemInternalId' keys, dates as ISO 8601 or as Excel serial numbers (DateTime Format setting).
  3. Builds the payload [{file, modified, rows}] like "Append to array variable", then performs the flow's Compose:
     base64 round trip of the template, replace __PA_REFRESHED__, replace __PA_DATA__ with string(payload) and '</' -> '<\\/'.
  4. Opens the output in headless Chrome and compares it with loading the same .xlsx files by hand in the
     normal single-file dashboard: same accidents, duplicates removed, FY and class counts. Also checks banner, footer,
     zero network requests, no console errors, a Summary containing '</script>' (escaped and NOT escaped),
     Save with data, Add more files, Clear, untouched template, broken JSON, flat-array and raw {value:[...]} payloads.

Run from the repo root (needs openpyxl + playwright with Chrome):  python3 tools/test-powerautomate.py
Writes only to a temp folder. Exit code 0 = all checks passed.
"""
import base64, datetime as dt, glob, json, os, random, re, subprocess, sys, tempfile, uuid
import openpyxl
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "downloads", "safety-dashboard-template.html")
SINGLE = os.path.join(ROOT, "downloads", "safety-dashboard.html")
CHROME = os.environ.get("CHROME", "/usr/bin/google-chrome")
TABLE = "MishapTracker"
DATE_COLS = {"Accident Date", "Due Date", "Compl. Date", "SUB DATE REPORTED"}
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
EVIL = "SM reported </script><script>window.__pwned=1</script> & </plaintext> \u201cquoted\u201d \u2014 test (fictional)."
STAMP = "30-Oct-26 1200"
FAILS = []
EVIL_ROW = 50   # a tracker record the roll-up does not repeat, so it is never replaced by de-duplication

def check(ok, label):
    print(("  PASS " if ok else "  FAIL ") + label)
    if not ok: FAILS.append(label)

def parse_text_date(s):
    m = re.match(r"^(\d{1,2})-([A-Za-z]{3})-(\d{2})$", str(s).strip())
    return dt.date(2000 + int(m.group(3)), MON.index(m.group(2).title()) + 1, int(m.group(1))) if m else None

def to_excel_online(path, date_mode):
    """Rows of the MishapTracker table as Excel Online returns them."""
    ws = openpyxl.load_workbook(path).active
    ref = ws.tables[TABLE].ref
    rows = list(ws[ref]); hdr = [c.value for c in rows[0]]; out = []
    for r in rows[1:]:
        o = {"@odata.etag": "", "ItemInternalId": str(uuid.uuid4())}
        for h, c in zip(hdr, r):
            v = c.value
            if v is None: s = ""
            elif isinstance(v, float) and v.is_integer(): s = str(int(v))
            else: s = str(v)
            d = parse_text_date(s) if h in DATE_COLS else None
            if d and date_mode == "iso": s = d.isoformat() + "T00:00:00.000Z"
            elif d and date_mode == "serial": s = str((d - dt.date(1899, 12, 30)).days)
            if h == "Time of Day" and re.match(r"^\d{4}$", s):
                hh, mm = int(s[:2]), int(s[2:])
                if date_mode == "iso": s = f"1899-12-30T{hh:02d}:{mm:02d}:00.000Z"
                elif date_mode == "serial": s = repr((hh * 60 + mm) / 1440)
            o[h] = s
        out.append(o)
    return out

def add_tables(folder):
    for p in glob.glob(os.path.join(folder, "*.xlsx")):
        st = os.stat(p); wb = openpyxl.load_workbook(p); ws = wb.active
        ws.add_table(Table(displayName=TABLE, ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"))
        wb.save(p); os.utime(p, (st.st_atime, st.st_mtime))       # keep the fixed file dates (newest file wins ties)

def payload_for(files, date_mode, evil_in=None):
    pl = []
    for p in files:
        rows = to_excel_online(p, date_mode)
        if evil_in and os.path.basename(p) == evil_in:
            rows[EVIL_ROW]["Summary"] = EVIL
        mod = dt.datetime.fromtimestamp(os.path.getmtime(p), dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        pl.append({"file": os.path.basename(p), "modified": mod, "rows": rows})
    random.Random(3).shuffle(pl)                                   # Get files order is not guaranteed
    return pl

def flow_output(payload, escape=True, raw=None):
    """The flow's Compose: replace(replace(base64ToString(template), TS, stamp), DATA, replace(string(payload), '</', '<\\/'))."""
    tmpl = base64.b64decode(base64.b64encode(open(TEMPLATE, "rb").read())).decode("utf-8")
    data = raw if raw is not None else json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if escape: data = data.replace("</", "<\\/")
    return tmpl.replace("__PA_" + "REFRESHED__", STAMP).replace("__PA_" + "DATA__", data)

def metrics(pg):
    pg.wait_for_function("!document.getElementById('loadBar').classList.contains('busy')", timeout=60000)
    m = {"bar": pg.inner_text("#lbText"), "fy": {}, "cls": {}}
    m["accidents"] = int(re.search(r"([\d,]+) accidents", m["bar"]).group(1).replace(",", ""))
    rep = pg.inner_text("#loadReport")
    d = re.search(r"(\d+) duplicate record", rep); m["dupes"] = int(d.group(1)) if d else 0
    for v in pg.eval_on_selector_all("#fFy option", "o=>o.map(x=>x.value)"):
        if not v: continue
        pg.select_option("#fFy", v); m["fy"]["FY" + v[2:]] = pg.inner_text("#asofLabel").split(" accident")[0]
    pg.select_option("#fFy", "")
    for v in pg.eval_on_selector_all("#fClass option", "o=>o.map(x=>x.value)"):
        if not v: continue
        pg.select_option("#fClass", v); m["cls"][v] = pg.inner_text("#asofLabel").split(" accident")[0]
    pg.select_option("#fClass", ""); pg.dispatch_event("#fClass", "change")
    return m

class Page:
    def __init__(self, browser, url):
        self.errors, self.net = [], []
        self.ctx = browser.new_context(accept_downloads=True, viewport={"width": 1440, "height": 900})
        self.pg = self.ctx.new_page()
        self.pg.on("console", lambda m: self.errors.append(m.text) if m.type == "error" else None)
        self.pg.on("pageerror", lambda e: self.errors.append(str(e)))
        self.pg.on("request", lambda r: self.net.append(r.url) if not r.url.startswith(("file:", "data:", "blob:")) else None)
        self.pg.goto(url); self.pg.wait_for_timeout(800)
    def close(self): self.ctx.close()

def write(tmp, name, html):
    p = os.path.join(tmp, name); open(p, "w", encoding="utf-8").write(html); return "file://" + p

def scenario(b, tmp, label, files, date_mode, escape, evil_file, baseline):
    print(f"\n== {label}: {len(files)} file(s), dates {date_mode}, flow escaping {'on' if escape else 'OFF'} ==")
    out = flow_output(payload_for(files, date_mode, evil_file), escape)
    check(("__PA_" + "DATA__") not in out and ("__PA_" + "REFRESHED__") not in out, "both placeholders replaced")
    P = Page(b, write(tmp, f"out-{label}.html", out)); pg = P.pg
    m = metrics(pg)
    print("   ", m["bar"], "| dupes", m["dupes"], "|", m["fy"])
    check(m["accidents"] == baseline["accidents"], f"accidents {m['accidents']} == manual upload {baseline['accidents']}")
    check(m["dupes"] == baseline["dupes"], f"duplicates removed {m['dupes']} == manual upload {baseline['dupes']}")
    check(m["fy"] == baseline["fy"], "per-FY counts match manual upload")
    check(m["cls"] == baseline["cls"], "per-class counts match manual upload")
    check(f"{len(files)} files" in m["bar"] or len(files) == 1, "one source per Excel file")
    ban = pg.inner_text("#realBanner") if pg.is_visible("#realBanner") else ""
    check(f"Data refreshed {STAMP} (automatic)" in ban, "banner: 'Data refreshed %s (automatic)'" % STAMP)
    check(pg.inner_text("#dataDisclaimer").startswith("Showing data loaded from your files"), "footer: loaded-data wording")
    check(pg.title().startswith("REAL DATA"), "tab title starts with REAL DATA")
    check(pg.evaluate("window.__pwned") is None, "Summary '</script><script>' did not execute")
    check(pg.is_enabled("#btnSave"), "Save with data enabled")
    if evil_file:
        pg.fill("#fSearch", "")
        rows = pg.locator("#tAll tbody tr[data-i]")
        found = False
        # open the record whose summary we replaced: search by its ASMIS # from the payload
        evil_asmis = [r for f in payload_for(files, date_mode, evil_file) if f["file"] == evil_file for r in f["rows"][EVIL_ROW:EVIL_ROW + 1]][0]["ASMIS #"]
        pg.fill("#fSearch", evil_asmis); pg.wait_for_timeout(300)
        if rows.count():
            rows.first.click(); pg.wait_for_timeout(300); body = pg.inner_text("#dBody"); found = "</script><script>window.__pwned=1</script>" in body and "\u201cquoted\u201d \u2014" in body
            pg.keyboard.press("Escape")
        pg.fill("#fSearch", "")
        check(found, "Summary with '</script>' and non-ASCII shown intact in the detail panel")
    return P, m

def main():
    tmp = tempfile.mkdtemp(prefix="pa-test-"); fy = os.path.join(tmp, "fy")
    subprocess.run([sys.executable, os.path.join(ROOT, "data-generation", "generate_asmis_fy_trackers.py"), "85", fy], check=True, stdout=subprocess.DEVNULL)
    add_tables(fy)
    files_all = sorted(glob.glob(os.path.join(fy, "*.xlsx")))
    files_two = [p for p in files_all if os.path.basename(p) in ("FY26 Tracker.xlsx", "FY26 Roll-up.xlsx")]
    roll = [p for p in files_all if "Roll-up" in p]
    print("temp:", tmp, "| files:", [os.path.basename(p) for p in files_all])
    tmpl = open(TEMPLATE, encoding="ascii").read()
    check(tmpl.count("__PA_" + "DATA__") == 1 and tmpl.count("__PA_" + "REFRESHED__") == 1, "template has each placeholder exactly once (ASCII file)")
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=CHROME, args=["--allow-file-access-from-files"])
        base = {}
        for key, files in (("all", files_all), ("two", files_two)):
            M = Page(b, "file://" + SINGLE); M.pg.set_input_files("#fileInput", files)
            M.pg.wait_for_function("/accidents/.test(document.getElementById('lbText').textContent) && !/Demo/.test(document.getElementById('lbText').textContent)", timeout=60000)
            base[key] = metrics(M.pg); M.close()
            print(f"manual upload baseline ({key}):", base[key]["bar"], "| dupes", base[key]["dupes"])
        check(base["all"]["accidents"] == 425 and base["all"]["dupes"] == 46, "baseline matches generator's expected 425 accidents / 46 duplicates")

        P, m = scenario(b, tmp, "iso-6files", files_all, "iso", True, "FY26 Tracker.xlsx", base["all"])
        check(not P.errors and not P.net, f"no console errors ({P.errors[:2]}) and zero network requests ({P.net[:2]})")
        pg = P.pg
        # Save with data from the automatic file
        with pg.expect_download() as dl: pg.click("#btnSave")
        sp = os.path.join(tmp, dl.value.suggested_filename); dl.value.save_as(sp); saved = open(sp, encoding="utf-8").read()
        check(re.match(r"SAFETY-REAL-DATA-\d\d[A-Z][a-z]{2}\d\d\.html$", os.path.basename(sp)) is not None, "Save with data file name " + os.path.basename(sp))
        check('<plaintext id="embedded' + '-data"' not in saved, "saved copy carries no Power Automate block (only the saved-data block)")
        S = Page(b, "file://" + sp); sm = metrics(S.pg)
        check(sm["accidents"] == base["all"]["accidents"] and "Data as of" in S.pg.inner_text("#realBanner"), f"saved copy reopens with {sm['accidents']} accidents and 'Data as of' banner")
        check(not S.errors and not S.net, "saved copy: no console errors, zero network requests"); S.close()
        # Add more files, then Clear
        pg.set_input_files("#addInput", roll); pg.wait_for_function("/7 files/.test(document.getElementById('lbText').textContent)", timeout=60000)
        am = metrics(pg)
        check(am["accidents"] == base["all"]["accidents"], f"Add more files (roll-up again): still {am['accidents']} accidents, {am['dupes']} duplicates removed")
        pg.click("#btnClear"); pg.wait_for_timeout(500); cm = metrics(pg)
        check(cm["accidents"] == base["all"]["accidents"] and "6 files" in cm["bar"], "Clear returns to the automatic data")
        check(not P.errors and not P.net, "still no console errors / network after save, add, clear"); P.close()

        P, _ = scenario(b, tmp, "serial-6files", files_all, "serial", True, None, base["all"])
        check(not P.errors and not P.net, "no console errors / network"); P.close()
        P, _ = scenario(b, tmp, "iso-2files-UNESCAPED", files_two, "iso", False, "FY26 Tracker.xlsx", base["two"])
        check(not P.errors and not P.net, f"no console errors / network even without the flow's '</' escape {P.errors[:2]}"); P.close()
        P, _ = scenario(b, tmp, "text-2files", files_two, "text", True, "FY26 Tracker.xlsx", base["two"])
        check(not P.errors and not P.net, "no console errors / network"); P.close()

        print("\n== other payload shapes ==")
        pl = payload_for(files_all, "iso")
        flat = [r for f in sorted(pl, key=lambda f: f["modified"]) for r in f["rows"]]
        P = Page(b, write(tmp, "flat.html", flow_output(None, True, json.dumps(flat)))); fm = metrics(P.pg)
        check(fm["accidents"] == base["all"]["accidents"], f"flat row array: {fm['accidents']} accidents ({fm['dupes']} duplicates removed)")
        check(not P.errors, "flat: no console errors"); P.close()
        one = [f for f in pl if f["file"] == "FY24 Tracker.xlsx"][0]
        P = Page(b, write(tmp, "raw.html", flow_output(None, True, json.dumps({"value": one["rows"]})))); rm = metrics(P.pg)
        check(rm["accidents"] == 85, f"raw Excel Online body {{value:[...]}}: {rm['accidents']} accidents")
        check(not P.errors, "raw body: no console errors"); P.close()

        print("\n== template untouched / broken ==")
        P = Page(b, "file://" + TEMPLATE); tm = metrics(P.pg)
        check("Demo data" in tm["bar"] and not P.pg.is_visible("#realBanner"), "untouched template shows demo data, no real-data banner")
        check(P.pg.inner_text("#dataDisclaimer").startswith("Data shown is fictional"), "untouched template: fictional footer")
        check(not P.errors and not P.net, "untouched template: no console errors, zero network requests"); P.close()
        P = Page(b, write(tmp, "broken.html", flow_output(None, False, '[{"file":"x.xlsx","rows":[{"Accident #":"1",'))); P.pg.wait_for_timeout(300)
        ban = P.pg.inner_text("#realBanner") if P.pg.is_visible("#realBanner") else ""
        check("could not be read" in ban and "FICTIONAL" in ban, "broken JSON: red banner says data could not be read, demo is labelled fictional")
        check(not P.errors, f"broken JSON: no uncaught errors {P.errors[:2]}"); P.close()
        b.close()
    print("\nRESULT:", "ALL PASSED" if not FAILS else f"{len(FAILS)} FAILED: {FAILS}")
    sys.exit(1 if FAILS else 0)

if __name__ == "__main__":
    main()
